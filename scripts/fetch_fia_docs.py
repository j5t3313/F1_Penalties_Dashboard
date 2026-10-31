import argparse
import hashlib
import json
import re
import sys
import time
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from data.loader import DRIVER_NAME_MAP, INVALID_SESSIONS

BASE_URL = "https://www.fia.com"
SEASON_URL = BASE_URL + "/documents/championships/fia-formula-one-world-championship-14/season/season-{year}-{season_id}"
AJAX_URL = BASE_URL + "/decision-document-list/ajax/{event_id}"
NOJS_URL = BASE_URL + "/decision-document-list/nojs/{event_id}"
SEASON_PATH_RE = re.compile(r"fia-formula-one-world-championship-14/season/season-(\d{4})-(\d+)$")
EVENT_ID_RE = re.compile(r"/decision-document-list/nojs/(\d+)")

KNOWN_SEASONS = {2025: 2071, 2026: 2072}
DEFAULT_YEARS = list(range(2020, 2027))

SOURCE_XLSX = ROOT / "data" / "F1Penalties.xlsx"
OUTPUT_XLSX = ROOT / "data" / "F1Penalties_with_docs.xlsx"
CACHE_DIR = ROOT / "data" / "fia_cache"
PROBE_DIR = CACHE_DIR / "probe"
UNMATCHED_ROWS_CSV = CACHE_DIR / "review_unmatched_rows.csv"
UNMATCHED_DOCS_CSV = CACHE_DIR / "review_unmatched_docs.csv"

USER_AGENT = "F1PenaltiesDashboard-DocBackfill/1.0 (personal research script; max 1 request per second)"
REQUEST_DELAY = 1.0
RETRY_STATUSES = (429, 503)
RETRY_BACKOFF = (30, 60, 120, 240)
LISTING_FINAL_AFTER = timedelta(days=3)
NUMBER_WORDS = {
    "one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7,
    "eight": 8, "nine": 9, "ten": 10, "fifteen": 15, "twenty": 20,
}
NUMBER = r"(\d+|" + "|".join(NUMBER_WORDS) + r")"

INCLUDE_KEYWORDS = ("infringement", "decision", "offence")
EXCLUDE_KEYWORDS = ("summons", "classification", "timetable", "notes", "recalled")

SESSION_MAP = {
    "free practice 1": "FP1",
    "free practice 2": "FP2",
    "free practice 3": "FP3",
    "practice 1": "FP1",
    "practice 2": "FP2",
    "practice 3": "FP3",
    "qualifying": "Q",
    "sprint qualifying": "SQ",
    "sprint shootout": "SQ",
    "sprint": "S",
    "race": "R",
    "reconnaissance laps": "Pre-Race",
    "pre race": "Pre-Race",
}

EVENT_RACE_ALIASES = {
    "abu dhabi grand prix": "Abu Dhabi",
    "australian grand prix": "Australia",
    "austrian grand prix": "Austria",
    "azerbaijan grand prix": "Azerbaijan",
    "bahrain grand prix": "Bahrain",
    "barcelona catalunya grand prix": "Spain",
    "belgian grand prix": "Belgium",
    "brazilian grand prix": "Brazil",
    "british grand prix": "Great Britain",
    "canadian grand prix": "Canada",
    "chinese grand prix": "China",
    "dutch grand prix": "Netherlands",
    "eifel grand prix": "Eifel",
    "emilia romagna grand prix": "Imola",
    "formula 1 70th anniversary grand prix": "Anniversary",
    "french grand prix": "France",
    "hungarian grand prix": "Hungary",
    "italian grand prix": "Italy",
    "japanese grand prix": "Japan",
    "las vegas grand prix": "Las Vegas",
    "mexican grand prix": "Mexico",
    "mexico city grand prix": "Mexico",
    "miami grand prix": "Miami",
    "monaco grand prix": "Monaco",
    "portuguese grand prix": "Portugal",
    "qatar grand prix": "Qatar",
    "russian grand prix": "Russia",
    "sakhir grand prix": "Sakhir",
    "sao paulo grand prix": "Brazil",
    "saudi arabian grand prix": "Saudi Arabia",
    "singapore grand prix": "Singapore",
    "spanish grand prix": "Spain",
    "styrian grand prix": "Styria",
    "turkish grand prix": "Turkey",
    "tuscan grand prix": "Tuscany",
    "united states grand prix": "United States",
}

EVENT_RACE_OVERRIDES = {
    2020: {
        "eifel grand prix": "France",
    },
    2021: {
        "styrian grand prix": "Austria",
    },
    2026: {
        "bahrain grand prix": "Malaysia*",
        "spanish grand prix": "Madrid",
        "barcelona catalunya grand prix": "Spain",
    },
}

TITLE_STOPWORDS = {
    "infringement", "decision", "offence", "corrected", "updated", "alleged", "car", "cars",
    "of", "the", "a", "an", "and", "to", "in", "on", "with", "by", "for", "at", "from", "during",
}

DASHES = dict.fromkeys(map(ord, "‐‑‒–—−"), "-")


def fold(value):
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9]+", " ", text.lower())
    return " ".join(text.split())


def clean_text(value):
    return unicodedata.normalize("NFKC", value or "").translate(DASHES)


def to_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return None if pd.isna(value) else float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def normalize_grid(value, outcome=None):
    outcome_text = fold(outcome)
    if "back of grid" in outcome_text or "back of the grid" in outcome_text:
        return "BACK"
    if "start from pit lane" in outcome_text:
        return "PIT"
    if value is None:
        return None
    number = to_number(value)
    if number is not None:
        return int(number)
    text = fold(value)
    if not text:
        return None
    if "pit" in text:
        return "PIT"
    if "back" in text:
        return "BACK"
    return text


def normalize_driver(name):
    if not name:
        return ""
    name = " ".join(str(name).split())
    return fold(DRIVER_NAME_MAP.get(name, name))


def drivers_match(a, b):
    tokens_a, tokens_b = set(a.split()), set(b.split())
    if not tokens_a or not tokens_b:
        return False
    return tokens_a <= tokens_b or tokens_b <= tokens_a


def keyword_tokens(text):
    tokens = set()
    for token in fold(text).split():
        if token in TITLE_STOPWORDS or token.isdigit() or len(token) < 3:
            continue
        tokens.add(token[:5])
    return tokens


class RateLimited(Exception):
    pass


class Fetcher:
    def __init__(self, delay=REQUEST_DELAY, cached_only=False):
        self.session = requests.Session()
        self.session.headers["User-Agent"] = USER_AGENT
        self.delay = delay
        self.cached_only = cached_only
        self.last_request = 0.0

    def request(self, method, url, **kwargs):
        kwargs.setdefault("timeout", 60)
        for attempt in range(len(RETRY_BACKOFF) + 1):
            wait = self.last_request + self.delay - time.monotonic()
            if wait > 0:
                time.sleep(wait)
            try:
                response = self.session.request(method, url, **kwargs)
            finally:
                self.last_request = time.monotonic()
            if response.status_code not in RETRY_STATUSES:
                return response
            if attempt == len(RETRY_BACKOFF):
                response.close()
                raise RateLimited(f"HTTP {response.status_code} after {len(RETRY_BACKOFF)} retries: {url}")
            retry_after = response.headers.get("Retry-After", "")
            pause = int(retry_after) if retry_after.isdigit() else RETRY_BACKOFF[attempt]
            print(f"  HTTP {response.status_code}, waiting {pause}s before retrying", flush=True)
            response.close()
            time.sleep(pause)
        return response

    def get(self, url, **kwargs):
        return self.request("GET", url, **kwargs)

    def head(self, url, **kwargs):
        return self.request("HEAD", url, **kwargs)


def season_url(year, season_id):
    return SEASON_URL.format(year=year, season_id=season_id)


def parse_season_options(html):
    soup = BeautifulSoup(html, "html.parser")
    found = {}
    for option in soup.find_all("option"):
        match = SEASON_PATH_RE.search(option.get("value", ""))
        if match:
            found[int(match.group(1))] = int(match.group(2))
    canonical = soup.find("link", rel="canonical")
    if canonical:
        match = SEASON_PATH_RE.search(canonical.get("href", ""))
        if match:
            found[int(match.group(1))] = int(match.group(2))
    return found


def resolve_season_ids(fetcher, years, cli_seasons):
    seasons = {}
    for year, season_id in KNOWN_SEASONS.items():
        response = fetcher.get(season_url(year, season_id))
        if response.status_code == 200:
            seasons.update(parse_season_options(response.text))
            break
    seasons.update(KNOWN_SEASONS)
    seasons.update(cli_seasons)
    missing = [year for year in years if year not in seasons]
    if missing:
        print(f"No season id found for {missing}. Pass them with --season YEAR=ID.")
    return {year: seasons[year] for year in years if year in seasons}


def list_events(html, year, season_id):
    soup = BeautifulSoup(html, "html.parser")
    event_ids = {}
    for link in soup.select("a.event-title"):
        match = EVENT_ID_RE.search(link.get("href", ""))
        if match:
            event_ids[" ".join(link.get_text(" ").split())] = int(match.group(1))
    marker = f"/season/season-{year}-{season_id}/event/"
    events = {}
    for option in soup.find_all("option"):
        value = option.get("value", "")
        if marker in value:
            name = " ".join(option.get_text(" ").split())
            events[name] = {"name": name, "url": urljoin(BASE_URL, value), "event_id": event_ids.get(name)}
    for name, event_id in event_ids.items():
        events.setdefault(name, {"name": name, "url": None, "event_id": event_id})
    return list(events.values())


def parse_document_rows(html):
    soup = BeautifulSoup(html, "html.parser")
    documents = []
    seen = set()
    for item in soup.select("li.document-row"):
        link = item.find("a", href=True)
        if not link or "decision-document/" not in link["href"]:
            continue
        href = link["href"]
        if href in seen:
            continue
        seen.add(href)
        title_el = item.select_one(".title")
        title = " ".join((title_el or link).get_text(" ").split())
        published_el = item.select_one(".published .date-display-single") or item.select_one(".published")
        published = " ".join(published_el.get_text(" ").split()) if published_el else ""
        recalled = item.select_one(".recalled-document") is not None or title.lower().startswith("recalled")
        listing_doc_no = None
        match = re.match(r"(?:Recalled\s*-\s*)?Doc\s+(\d+)\s*-\s*(.*)$", title, re.I)
        if match:
            listing_doc_no = int(match.group(1))
            title = match.group(2).strip()
        elif title.lower().startswith("recalled"):
            title = re.sub(r"^recalled\s*-\s*", "", title, flags=re.I)
        documents.append({
            "href": href,
            "url": urljoin(BASE_URL, href),
            "title": title,
            "published": published,
            "recalled": recalled,
            "listing_doc_no": listing_doc_no,
        })
    return documents


def ajax_html(text):
    try:
        commands = json.loads(text)
    except ValueError:
        return ""
    return "".join(
        command.get("data", "")
        for command in commands
        if isinstance(command, dict) and command.get("command") == "insert" and isinstance(command.get("data"), str)
    )


def fetch_event_documents(fetcher, event):
    if event["url"]:
        response = fetcher.get(event["url"])
        if response.status_code == 200:
            documents = parse_document_rows(response.text)
            if documents:
                return documents
    if event["event_id"]:
        response = fetcher.get(AJAX_URL.format(event_id=event["event_id"]))
        if response.status_code == 200:
            return parse_document_rows(ajax_html(response.text))
    return []


def listing_path(year, event):
    slug = re.sub(r"[^a-z0-9]+", "_", fold(event["name"])).strip("_")
    return CACHE_DIR / str(year) / "listings" / f"{slug}.json"


def listing_is_final(documents):
    published = []
    for document in documents:
        match = re.search(r"(\d{2})\.(\d{2})\.(\d{2})\s+(\d{2}):(\d{2})", document.get("published", ""))
        if match:
            day, month, year, hour, minute = map(int, match.groups())
            published.append(datetime(2000 + year, month, day, hour, minute))
    return bool(published) and datetime.now() - max(published) > LISTING_FINAL_AFTER


def event_documents(fetcher, year, event, refresh_listings):
    path = listing_path(year, event)
    if not refresh_listings and path.exists():
        saved = json.loads(path.read_text(encoding="utf-8"))
        if saved.get("final"):
            return saved["documents"]
    documents = fetch_event_documents(fetcher, event)
    if documents:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps({
            "event": event["name"],
            "saved_at": datetime.now().isoformat(timespec="seconds"),
            "final": listing_is_final(documents),
            "documents": documents,
        }, indent=1), encoding="utf-8")
    return documents


def keep_document(document):
    title = document["title"].lower()
    if document["recalled"]:
        return False
    if any(keyword in title for keyword in EXCLUDE_KEYWORDS):
        return False
    return any(keyword in title for keyword in INCLUDE_KEYWORDS)


def mark_superseded(documents):
    def base_title(title):
        return fold(re.sub(r"\b(corrected|updated)\b", "", title, flags=re.I))

    for document in documents:
        if not re.search(r"\b(corrected|updated)\b", document["title"], re.I):
            continue
        for other in documents:
            if other is document or other["event"] != document["event"]:
                continue
            if base_title(other["title"]) == base_title(document["title"]) and not re.search(r"\b(corrected|updated)\b", other["title"], re.I):
                other["status"] = "superseded"


def cache_path(year, url):
    name = unquote(urlparse(url).path.rsplit("/", 1)[-1])
    stem = re.sub(r'[<>:"/\\|?*\s]+', "_", name[:-4] if name.lower().endswith(".pdf") else name).strip("._")
    digest = hashlib.md5(url.encode("utf-8")).hexdigest()[:8]
    return CACHE_DIR / str(year) / f"{stem[:60]}_{digest}.pdf"


def download_pdf(fetcher, url, path):
    if path.exists() and path.stat().st_size > 0:
        return None
    response = fetcher.get(url)
    if response.status_code != 200:
        return f"http_{response.status_code}"
    if not response.content.startswith(b"%PDF"):
        return "not_pdf"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(response.content)
    return None


def pdf_text(path):
    reader = PdfReader(str(path))
    return clean_text("\n".join(page.extract_text() or "" for page in reader.pages))


def word_number(value):
    return int(value) if value.isdigit() else NUMBER_WORDS[value.lower()]


def first_match(patterns, text):
    for pattern in patterns:
        match = re.search(pattern, text, re.I)
        if match:
            return match
    return None


def parse_pdf(text):
    lines = [" ".join(line.split()) for line in text.splitlines()]
    fields = {
        "header_year": None, "header_event": None, "pdf_doc_no": None, "car": None,
        "driver_raw": None, "session_raw": None, "decision": "",
        "fine": None, "grid": None, "time": None, "points": None,
    }
    for line in lines:
        match = re.match(r"^(20\d{2})\s+(.*\bGRAND PRIX\b.*)$", line, re.I)
        if match:
            fields["header_year"] = int(match.group(1))
            fields["header_event"] = match.group(2).strip()
            break
    match = re.search(r"^\s*Document\s+(\d+)", text, re.I | re.M)
    if match:
        fields["pdf_doc_no"] = int(match.group(1))
    match = re.search(r"No\s*/\s*Driver\s+(\d+)\s*-\s*(.+)", text, re.I)
    if match:
        fields["car"] = int(match.group(1))
        fields["driver_raw"] = " ".join(match.group(2).split())
    header_block = re.split(r"^\s*Fact\b", text, maxsplit=1, flags=re.M)[0]
    match = re.search(r"^\s*Session\s+(.+)$", header_block, re.M)
    if match:
        fields["session_raw"] = " ".join(match.group(1).split())
    match = re.search(
        r"^\s*Decision\s+(.*?)(?=^\s*Reason\b|^\s*Competitors are reminded|\Z)",
        text, re.I | re.M | re.S,
    )
    decision = " ".join(match.group(1).split()) if match else ""
    fields["decision"] = decision

    amount = r"(\d{1,3}(?:[.,  ]\d{3})+|\d+)"
    match = first_match([rf"(?:€|EUR|euros?)\s*{amount}", rf"{amount}\s*(?:€|EUR|euros?)\b"], decision)
    if match:
        fields["fine"] = float(re.sub(r"\D", "", match.group(1)))

    match = first_match([
        rf"drop of {NUMBER} grid (?:positions?|places?)",
        rf"\b{NUMBER}[\s-]*(?:places?|positions?)\s+grid\s+(?:penalty|drop)",
        rf"\b{NUMBER}\s+grid\s+(?:places?|positions?)\s+(?:penalty|drop)",
    ], decision)
    if match:
        fields["grid"] = word_number(match.group(1))
    elif re.search(r"back of the (?:starting )?grid", decision, re.I):
        fields["grid"] = "BACK"
    elif re.search(r"start (?:the (?:race|sprint) )?from the pit ?lane", decision, re.I):
        fields["grid"] = "PIT"

    match = re.search(rf"\b{NUMBER}\s*-?\s*seconds?\s+(?:time\s+)?penalty", decision, re.I)
    if match:
        fields["time"] = float(word_number(match.group(1)))

    match = re.search(rf"\b{NUMBER}\s+penalty\s+points?", decision, re.I)
    if match:
        fields["points"] = float(word_number(match.group(1)))
    return fields


def map_session(raw, unmapped):
    if not raw:
        return None
    key = fold(raw)
    if key in SESSION_MAP:
        return SESSION_MAP[key]
    unmapped.add(raw)
    return None


def resolve_race(year, event_name, sheet_races):
    key = fold(event_name)
    if re.search(r"\btests?\b|\btesting\b", key):
        return None
    candidate =EVENT_RACE_OVERRIDES.get(year, {}).get(key) or EVENT_RACE_ALIASES.get(key)
    if candidate:
        return candidate if candidate in sheet_races else None
    hits = [race for race in sheet_races if fold(race) and fold(race) in key]
    return hits[0] if len(hits) == 1 else None


def print_race_table(year, events, sheet_races):
    print(f"\n{year}: FIA event -> Race")
    used = set()
    for event in sorted(events, key=lambda e: e["name"]):
        race = event["race"]
        used.add(race)
        print(f"  {event['name']:<45} -> {race if race else '(no match)'}")
    missing = sorted(set(sheet_races) - used)
    if missing:
        print(f"  Sheet races with no FIA event: {', '.join(missing)}")


def load_rows(path, years):
    workbook = load_workbook(path, data_only=True, read_only=True)
    rows = []
    for sheet in workbook.worksheets:
        if not sheet.title.isdigit() or int(sheet.title) not in years:
            continue
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            continue
        columns = {str(name).strip(): index for index, name in enumerate(header) if name is not None}

        def cell(values, name):
            index = columns.get(name)
            if index is None or index >= len(values):
                return None
            value = values[index]
            if isinstance(value, str):
                value = value.strip()
                return value or None
            return value

        for excel_row, values in enumerate(iterator, start=2):
            if all(value is None for value in values):
                continue
            year = to_number(cell(values, "Year")) or int(sheet.title)
            session = cell(values, "Session")
            outcome = cell(values, "Outcome")
            rows.append({
                "sheet": sheet.title,
                "excel_row": excel_row,
                "Year": int(year),
                "Round": cell(values, "Round"),
                "Race": cell(values, "Race"),
                "Driver": cell(values, "Driver"),
                "Session": session,
                "Allegation": cell(values, "Allegation"),
                "Allegation_Raw": cell(values, "Allegation_Raw"),
                "Outcome": outcome,
                "Fine": cell(values, "Fine"),
                "Grid Penalty": cell(values, "Grid Penalty"),
                "Time Penalty (in seconds)": cell(values, "Time Penalty (in seconds)"),
                "Penalty Points": cell(values, "Penalty Points"),
                "Doc_URL": cell(values, "Doc_URL"),
                "driver_key": normalize_driver(cell(values, "Driver")),
                "session_key": None if session is None or str(session) in INVALID_SESSIONS else str(session),
                "fine_key": to_number(cell(values, "Fine")),
                "grid_key": normalize_grid(cell(values, "Grid Penalty"), outcome),
                "time_key": to_number(cell(values, "Time Penalty (in seconds)")),
                "points_key": to_number(cell(values, "Penalty Points")),
                "keywords": keyword_tokens(f"{cell(values, 'Allegation') or ''} {cell(values, 'Allegation_Raw') or ''}"),
            })
    workbook.close()
    return rows


def penalty_score(document, row):
    score = 0.0
    conflicts = []
    for name, doc_key, row_key in [
        ("Fine", "fine", "fine_key"),
        ("Grid", "grid", "grid_key"),
        ("Time", "time", "time_key"),
        ("Points", "points", "points_key"),
    ]:
        doc_value, row_value = document.get(doc_key), row.get(row_key)
        if doc_value is None and (row_value is None or row_value == 0):
            continue
        if doc_value is None or row_value is None:
            score -= 0.25
        elif doc_value == row_value:
            score += 1
        else:
            conflicts.append(f"{name} doc={doc_value} row={row_value}")
    return score, conflicts


def candidate_rows(document, rows):
    return [
        row for row in rows
        if row["Year"] == document["year"]
        and row["Race"] == document["race"]
        and drivers_match(row["driver_key"], document["driver_key"])
    ]


def session_compatible(document, row):
    return document["session"] is None or row["session_key"] is None or row["session_key"] == document["session"]


def match_documents(documents, rows):
    pairs = []
    for document in documents:
        if document["status"] != "parsed":
            continue
        for row in candidate_rows(document, rows):
            if row["Doc_URL"] or not session_compatible(document, row):
                continue
            score, conflicts = penalty_score(document, row)
            if conflicts:
                continue
            keywords = len(document["keywords"] & row["keywords"])
            pairs.append(((-score, -keywords, row["sheet"], row["excel_row"], document["doc_no"] or 0), document, row))
    pairs.sort(key=lambda pair: pair[0])
    assigned_rows, assigned_docs = {}, set()
    for _, document, row in pairs:
        row_id = (row["sheet"], row["excel_row"])
        if row_id in assigned_rows or id(document) in assigned_docs:
            continue
        assigned_rows[row_id] = document
        assigned_docs.add(id(document))
        document["status"] = "matched"
        document["matched_row"] = f"{row['sheet']}!{row['excel_row']}"
    return assigned_rows


def describe_candidates(row, documents, assigned_rows):
    notes = []
    for document in documents:
        if document["year"] != row["Year"] or document["race"] != row["Race"]:
            continue
        if not document.get("driver_key") or not drivers_match(row["driver_key"], document["driver_key"]):
            continue
        reasons = []
        if document["status"] == "matched":
            reasons.append(f"assigned to {document['matched_row']}")
        elif document["status"] != "parsed":
            reasons.append(document["status"])
        if not session_compatible(document, row):
            reasons.append(f"session doc={document['session']} row={row['session_key']}")
        _, conflicts = penalty_score(document, row)
        reasons.extend(conflicts)
        notes.append(f"Doc {document['doc_no']} | {document['title']} | {'; '.join(reasons) or 'ok'} | {document['url']}")
    return " || ".join(notes)


def process_documents(fetcher, year, year_documents, linked_urls, unmapped_sessions):
    mark_superseded(year_documents)
    for number, document in enumerate(year_documents, start=1):
        print(f"  [{number}/{len(year_documents)}] {document['event']}: {document['title']}")
        document.setdefault("driver_key", "")
        document.setdefault("session", None)
        document.setdefault("keywords", keyword_tokens(document["title"]))
        document.setdefault("doc_no", document["listing_doc_no"])
        if document["url"] in linked_urls:
            document["status"] = "already_linked"
            continue
        if document["status"] == "superseded":
            continue
        if not document["race"]:
            document["status"] = "race_unmapped"
            continue
        path = cache_path(year, document["url"])
        if fetcher.cached_only and not path.exists():
            document["status"] = "not_downloaded"
            continue
        error = download_pdf(fetcher, document["url"], path)
        if error:
            document["status"] = f"download_failed:{error}"
            continue
        try:
            fields = parse_pdf(pdf_text(path))
        except Exception as exc:
            document["status"] = f"parse_error:{type(exc).__name__}"
            continue
        document.update(fields)
        document["doc_no"] = fields["pdf_doc_no"] or document["listing_doc_no"]
        document["session"] = map_session(fields["session_raw"], unmapped_sessions)
        document["driver_key"] = normalize_driver(fields["driver_raw"])
        if fields["header_year"] and fields["header_year"] != year:
            document["status"] = f"year_mismatch:{fields['header_year']}"
        elif not document["race"]:
            document["status"] = "race_unmapped"
        elif not document["driver_key"]:
            document["status"] = "no_driver"
        else:
            document["status"] = "parsed"


def crawl(fetcher, years, cli_seasons, rows, map_only, refresh_listings):
    seasons = resolve_season_ids(fetcher, years, cli_seasons)
    sheet_races = {}
    for row in rows:
        if row["Race"]:
            sheet_races.setdefault(row["Year"], set()).add(row["Race"])
    linked_urls = {row["Doc_URL"] for row in rows if row["Doc_URL"]}

    season_events = {}
    for year, season_id in seasons.items():
        response = fetcher.get(season_url(year, season_id))
        if response.status_code != 200:
            print(f"{year}: season page returned {response.status_code}")
            continue
        events = list_events(response.text, year, season_id)
        for event in events:
            event["race"] = resolve_race(year, event["name"], sheet_races.get(year, set()))
        season_events[year] = events
        print_race_table(year, events, sheet_races.get(year, set()))

    if map_only:
        return None

    documents = []
    unmapped_sessions = set()
    for year, events in season_events.items():
        year_documents = []
        for event in events:
            listed = [document for document in event_documents(fetcher, year, event, refresh_listings) if keep_document(document)]
            print(f"{year} {event['name']}: {len(listed)} decision documents")
            for document in listed:
                document.update({"year": year, "event": event["name"], "race": event["race"], "status": "listed"})
            year_documents.extend(listed)
        process_documents(fetcher, year, year_documents, linked_urls, unmapped_sessions)
        documents.extend(year_documents)

    if unmapped_sessions:
        print("\nUnmapped session values (session filter ignored for these):")
        for value in sorted(unmapped_sessions):
            print(f"  {value}")
    return documents


def write_outputs(rows, documents, assigned_rows, years):
    workbook = load_workbook(SOURCE_XLSX)
    written = 0
    for (sheet_name, excel_row), document in assigned_rows.items():
        sheet = workbook[sheet_name]
        header = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value is not None}
        column = header.get("Doc_URL")
        if column is None:
            column = sheet.max_column + 1
            sheet.cell(row=1, column=column, value="Doc_URL")
        target = sheet.cell(row=excel_row, column=column)
        if target.value not in (None, ""):
            continue
        target.value = document["url"]
        written += 1
    workbook.save(OUTPUT_XLSX)
    print(f"\nWrote {written} links to {OUTPUT_XLSX}")

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    row_columns = [
        "sheet", "excel_row", "Year", "Round", "Race", "Driver", "Session", "Allegation", "Outcome",
        "Fine", "Grid Penalty", "Time Penalty (in seconds)", "Penalty Points",
    ]
    unmatched_rows = []
    for row in rows:
        if row["Doc_URL"] or (row["sheet"], row["excel_row"]) in assigned_rows:
            continue
        record = {column: row[column] for column in row_columns}
        record["candidate_docs"] = describe_candidates(row, documents, assigned_rows)
        unmatched_rows.append(record)
    pd.DataFrame(unmatched_rows, columns=row_columns + ["candidate_docs"]).to_csv(UNMATCHED_ROWS_CSV, index=False, encoding="utf-8-sig")

    doc_columns = [
        "year", "event", "race", "status", "doc_no", "title", "published", "url",
        "header_year", "header_event", "car", "driver_raw", "session_raw", "session",
        "fine", "grid", "time", "points", "decision",
    ]
    unmatched_docs = [
        {column: document.get(column) for column in doc_columns}
        for document in documents
        if document["status"] not in ("matched", "already_linked")
    ]
    pd.DataFrame(unmatched_docs, columns=doc_columns).to_csv(UNMATCHED_DOCS_CSV, index=False, encoding="utf-8-sig")
    print(f"Wrote {UNMATCHED_ROWS_CSV}")
    print(f"Wrote {UNMATCHED_DOCS_CSV}")

    print("\nSeason summary")
    print(f"  {'Year':<6}{'Rows matched':>14}{'Rows unmatched':>16}{'Docs unmatched':>16}")
    for year in years:
        matched = sum(1 for (sheet_name, _), document in assigned_rows.items() if document["year"] == year)
        unmatched = sum(1 for record in unmatched_rows if record["Year"] == year)
        docs_unmatched = sum(1 for document in unmatched_docs if document["year"] == year)
        print(f"  {year:<6}{matched:>14}{unmatched:>16}{docs_unmatched:>16}")


def run_backfill(years, cli_seasons, map_only, delay, cached_only, refresh_listings):
    rows = load_rows(SOURCE_XLSX, set(years))
    fetcher = Fetcher(delay, cached_only)
    documents = crawl(fetcher, years, cli_seasons, rows, map_only, refresh_listings)
    if documents is None:
        return
    assigned_rows = match_documents(documents, rows)
    write_outputs(rows, documents, assigned_rows, years)


def run_probe(years, cli_seasons, delay):
    PROBE_DIR.mkdir(parents=True, exist_ok=True)
    fetcher = Fetcher(delay)
    year = years[-1] if len(years) == 1 else 2025
    seasons = resolve_season_ids(fetcher, [year], cli_seasons)
    if year not in seasons:
        return
    season_id = seasons[year]
    response = fetcher.get(season_url(year, season_id))
    season_file = PROBE_DIR / f"season_{year}.html"
    season_file.write_text(response.text, encoding="utf-8")
    print(f"Season page {response.status_code}, saved {season_file}")
    print(f"Seasons in selector: {parse_season_options(response.text)}")
    events = list_events(response.text, year, season_id)
    print(f"Events found: {len(events)}")
    for event in events:
        print(f"  {event['name']} | id={event['event_id']} | {event['url']}")
    if not events:
        return

    event = next((e for e in events if e["url"] and e["event_id"]), events[0])
    slug = re.sub(r"[^a-z0-9]+", "_", event["name"].lower()).strip("_")
    if event["url"]:
        response = fetcher.get(event["url"])
        path = PROBE_DIR / f"event_{slug}.html"
        path.write_text(response.text, encoding="utf-8")
        print(f"Event page {response.status_code}, {len(parse_document_rows(response.text))} documents, saved {path}")
    if event["event_id"]:
        response = fetcher.get(NOJS_URL.format(event_id=event["event_id"]), allow_redirects=False)
        path = PROBE_DIR / f"event_{event['event_id']}_nojs.html"
        path.write_text(response.text, encoding="utf-8")
        print(f"nojs page {response.status_code} location={response.headers.get('Location')}, saved {path}")
        response = fetcher.get(AJAX_URL.format(event_id=event["event_id"]))
        path = PROBE_DIR / f"event_{event['event_id']}_ajax.json"
        path.write_text(response.text, encoding="utf-8")
        print(f"ajax endpoint {response.status_code}, {len(parse_document_rows(ajax_html(response.text)))} documents, saved {path}")

    documents = fetch_event_documents(fetcher, event)
    kept = [document for document in documents if keep_document(document)]
    print(f"\n{event['name']}: {len(documents)} documents listed, {len(kept)} kept")
    for document in documents:
        flag = "KEEP" if document in kept else "skip"
        print(f"  [{flag}] Doc {document['listing_doc_no']} - {document['title']} ({document['published']})")
    for document in kept[:3]:
        path = PROBE_DIR / Path(cache_path(year, document["url"])).name
        error = download_pdf(fetcher, document["url"], path)
        if error:
            print(f"\n{document['url']}: {error}")
            continue
        fields = parse_pdf(pdf_text(path))
        print(f"\n{document['title']}")
        for key, value in fields.items():
            print(f"  {key}: {value}")
        print(f"  session -> {map_session(fields['session_raw'], set())}")
        print(f"  driver -> {normalize_driver(fields['driver_raw'])}")


def run_check_urls(path, delay):
    fetcher = Fetcher(delay)
    workbook = load_workbook(path, data_only=True, read_only=True)
    entries = []
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            continue
        columns = {str(name).strip(): index for index, name in enumerate(header) if name is not None}
        if "Doc_URL" not in columns:
            continue

        def value(values, name):
            i = columns.get(name)
            return values[i] if i is not None and i < len(values) else None

        for excel_row, values in enumerate(iterator, start=2):
            url = value(values, "Doc_URL")
            if not isinstance(url, str) or not url.strip():
                continue
            entries.append((sheet.title, excel_row, value(values, "Race"), value(values, "Driver"), value(values, "Session"), url.strip()))
    workbook.close()

    failures = []
    checked = 0
    for number, (sheet_name, excel_row, race, driver, session, url) in enumerate(entries, start=1):
        try:
            response = fetcher.head(url, allow_redirects=True)
            if response.status_code != 200:
                response = fetcher.get(url, allow_redirects=True, stream=True)
                response.close()
            status = response.status_code
            content_type = response.headers.get("Content-Type", "")
        except requests.RequestException as exc:
            status, content_type = type(exc).__name__, ""
        except RateLimited as exc:
            print(f"Stopped: the FIA site is rate-limiting requests ({exc}). Try again later.")
            break
        checked += 1
        print(f"  [{number}/{len(entries)}] {status} {sheet_name}!{excel_row}")
        if status != 200 or "pdf" not in content_type.lower():
            failures.append((sheet_name, excel_row, race, driver, session, status, content_type, url))
    print(f"Checked {checked} of {len(entries)} URLs in {path}")
    if not failures:
        print("All returned 200 with a PDF.")
        return
    print(f"{len(failures)} did not return 200 with a PDF:")
    for sheet_name, excel_row, race, driver, session, status, content_type, url in failures:
        print(f"  {sheet_name}!{excel_row} {race} | {driver} | {session} | {status} {content_type}")
        print(f"    {url}")


def parse_seasons(values):
    seasons = {}
    for value in values or []:
        year, _, season_id = value.partition("=")
        if not year.isdigit() or not season_id.isdigit():
            raise argparse.ArgumentTypeError(f"Bad --season value: {value}")
        seasons[int(year)] = int(season_id)
    return seasons


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)
    parser = argparse.ArgumentParser(description="Link F1Penalties.xlsx rows to FIA decision documents.")
    parser.add_argument("--years", nargs="+", type=int, default=DEFAULT_YEARS)
    parser.add_argument("--season", action="append", metavar="YEAR=ID")
    parser.add_argument("--probe", action="store_true")
    parser.add_argument("--map-only", action="store_true")
    parser.add_argument("--check-urls", nargs="?", const=str(SOURCE_XLSX), metavar="WORKBOOK")
    parser.add_argument("--delay", type=float, default=REQUEST_DELAY)
    parser.add_argument("--cached-only", action="store_true")
    parser.add_argument("--refresh-listings", action="store_true")
    args = parser.parse_args()
    cli_seasons = parse_seasons(args.season)
    delay = max(args.delay, REQUEST_DELAY)

    try:
        if args.check_urls:
            run_check_urls(Path(args.check_urls), delay)
        elif args.probe:
            run_probe(args.years, cli_seasons, delay)
        else:
            run_backfill(sorted(set(args.years)), cli_seasons, args.map_only, delay, args.cached_only, args.refresh_listings)
    except RateLimited as exc:
        print(f"\nStopped: the FIA site is rate-limiting requests ({exc}).")
        print("Downloaded PDFs are cached in data/fia_cache, so rerunning later resumes where this stopped.")
        sys.exit(1)


if __name__ == "__main__":
    main()
